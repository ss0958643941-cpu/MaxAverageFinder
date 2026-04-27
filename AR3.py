import os
import sys
import pandas as pd
from tqdm import tqdm
import re
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side

def get_natural_sort_key(s):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', str(s))]

def get_date_from_folder(folder_name):
    """Extracts date from folder name, e.g., 20251117-13.00 -> 20251117"""
    match = re.search(r'\d{8}', folder_name)
    return match.group(0) if match else "Unknown_Date"

def get_base_sample_name(sample_id):
    """ตัดตัวเลขด้านหลังออกเพื่อหากลุ่มหลัก เช่น S_01 -> S"""
    return re.split(r'[_ -]', sample_id)[0]

# ==========================================
# Terminal Interface & Deep Scan Module
# ==========================================
def scan_and_display_file_statistics(root_folder):
    print("\n[ SYSTEM ] Initiating directory scan...")
    max_lines_overall = 0
    file_stats = []
    time_folders_dict = {}
    
    for root, dirs, files in os.walk(root_folder):
        if "pipeline_output" in root.lower() or "test" in root.lower(): 
            continue
            
        valid_files = [f for f in files if f.endswith(('.csv', '.txt')) 
                       and "air" not in f.lower() 
                       and "test" not in f.lower()]
        
        if valid_files:
            time_folders_dict[root] = valid_files

    if not time_folders_dict:
        return 0

    sorted_folder_paths = sorted(time_folders_dict.keys(), key=get_natural_sort_key)

    for folder_path in sorted_folder_paths:
        time_folder_name = os.path.basename(folder_path)
        valid_files = sorted(time_folders_dict[folder_path], key=get_natural_sort_key)
        
        for f_name in valid_files:
            file_path = os.path.join(folder_path, f_name)
            sample_name = os.path.splitext(f_name)[0].upper() 
            
            try:
                lines = 0
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                    for line in file:
                        lines += 1
                        if '[EndOfFile]' in line: break
                            
                if lines > max_lines_overall:
                    max_lines_overall = lines
                    
                if lines > 0:
                    file_stats.append(f"  |-- [FOLDER: {time_folder_name:<25}] Sample: {sample_name:<8} -> {lines} Rows")
            except: continue

    if file_stats:
        print("[ INFO ] File structure analysis complete:")
        for stat in file_stats:
            print(stat)
            
    return max_lines_overall

def get_user_target_row(max_lines):
    print("-" * 70)
    print(f"[ INFO ] Scan complete. Maximum row count is {max_lines}.")
    print("[ INFO ] The program will extract data exactly at the row number you input.")
    print("[ INFO ] Press [ENTER] without typing anything to cancel and exit.")
    
    while True:
        raw_input = input(f"[ INPUT ] Enter the TARGET row (1 - {max_lines}): ").strip()
        
        if not raw_input:
            print("[ CANCEL ] Operation aborted. Exiting...")
            sys.exit(0)

        try:
            target_row = int(raw_input) 
            if 0 < target_row <= max_lines:
                break
            print(f"[ ERROR ] Out of range. Please enter 1 - {max_lines}.")
        except ValueError:
            print("[ ERROR ] Please enter a valid integer.")

    return target_row

# ==========================================
# Data Extraction Module
# ==========================================
def extract_value_at_target_row(file_path, target_row):
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            for line_idx, line in enumerate(f, start=1):
                if line_idx == target_row:
                    line = line.strip()
                    parts = line.replace(';', ' ').replace(',', ' ').split()
                    if len(parts) >= 2:
                        try:
                            return float(parts[1]), float(parts[0]) 
                        except ValueError: return None, None
                    break
        return None, None
    except: return None, None

# ==========================================
# Output Generation Module (Stats at Bottom Rows)
# ==========================================
def save_professional_excel(df_data, file_name):
    if df_data.empty: return

    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(base_dir, file_name)

    try:
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            params = [('Intensity', 'Parameter: Max Intensity'), 
                      ('Time / Wavelength', 'Parameter: Time / Wavelength')]
            
            pastel_colors = ["DDEBF7", "E2EFDA", "FFF2CC", "FCE4D6", "E4DFEC", "D9D2E9", "F4CCCC"]
            
            for param_val, param_title in params:
                df_param = df_data[df_data['Parameter'] == param_val]
                if df_param.empty: continue

                df_pivot = df_param.pivot_table(
                    index=['Date', 'Time Folders'],
                    columns='Sample ID',
                    values='Measured Value',
                    aggfunc='first'
                ).reset_index()

                df_pivot.columns.name = None
                
                # 🔥 เพิ่ม 3 แถวสรุป (Average, Min, SD) ไว้ล่างสุดในแนวตั้ง 🔥
                numeric_cols = df_pivot.columns.drop(['Date', 'Time Folders'])
                
                avg_row = {'Date': '', 'Time Folders': 'Average'}
                min_row = {'Date': '', 'Time Folders': 'Minimum'}
                sd_row = {'Date': '', 'Time Folders': 'SD'}
                
                for col in numeric_cols:
                    # คำนวณแบบไม่เอาค่าว่าง (NaN) มารวม
                    avg_row[col] = df_pivot[col].mean()
                    min_row[col] = df_pivot[col].min()
                    sd_row[col] = df_pivot[col].std()
                
                # นำแถวสถิติทั้ง 3 ต่อท้ายตารางหลัก
                stat_df = pd.DataFrame([avg_row, min_row, sd_row])
                df_pivot = pd.concat([df_pivot, stat_df], ignore_index=True)
                
                # แปลง NaN ให้เป็น None เพื่อให้เซลล์ว่างสะอาดตา
                df_pivot = df_pivot.where(pd.notnull(df_pivot), None)
                
                sheet_name = param_val.replace(' / ', '_')[:30]
                df_pivot.to_excel(writer, sheet_name=sheet_name, startrow=1, index=False)
                ws = writer.sheets[sheet_name]
                
                ws['A1'] = param_title
                ws['A1'].font = Font(bold=True, color="002060") 
                
                header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")
                stat_fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid") # สีเทาสำหรับแถวสรุปล่างสุด
                index_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
                thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                     top=Side(style='thin'), bottom=Side(style='thin'))
                
                # สร้าง Map สีเหมือนเดิม
                unique_bases = []
                for c in numeric_cols:
                    base = get_base_sample_name(c)
                    if base not in unique_bases: unique_bases.append(base)
                base_color_map = {b: pastel_colors[i % len(pastel_colors)] for i, b in enumerate(unique_bases)}
                
                max_row = ws.max_row
                
                for i, col in enumerate(ws.columns, 1):
                    max_len = 0
                    col_name = df_pivot.columns[i - 1]
                    
                    if col_name in ['Date', 'Time Folders']:
                        col_fill = index_fill
                    else:
                        base_name = get_base_sample_name(col_name)
                        col_fill = PatternFill(start_color=base_color_map[base_name], 
                                               end_color=base_color_map[base_name], fill_type="solid")
                    
                    for cell in col:
                        cell.border = thin_border
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        
                        if cell.row == 2:
                            cell.fill = header_fill
                            cell.font = header_font
                        # ถ้าเป็น 3 บรรทัดสุดท้าย (Average, Min, SD) ให้ระบายสีเทาตัวหนา
                        elif cell.row >= max_row - 2:
                            cell.fill = stat_fill
                            cell.font = Font(bold=True)
                        else:
                            cell.fill = col_fill
                            
                        if cell.value is not None:
                            max_len = max(max_len, len(str(cell.value)))
                            
                    ws.column_dimensions[col[0].column_letter].width = max_len + 3

        print(f"[ OK ] Export successful: {file_name}")
    except PermissionError:
        print(f"[ ERROR ] Permission denied. Please close the file {file_name} and try again.")

def process_professional_layout(root_folder, target_row):
    print(f"\n[ EXECUTE ] Extracting values at row {target_row}...")
    data_list = []
    
    for root, dirs, files in os.walk(root_folder):
        if "pipeline_output" in root.lower() or "test" in root.lower(): continue
        
        valid_files = [f for f in files if f.endswith(('.csv', '.txt')) 
                       and "air" not in f.lower() 
                       and "test" not in f.lower()]
        
        if valid_files:
            folder_name = os.path.basename(root)
            date_val = get_date_from_folder(folder_name)
            
            for f_name in sorted(valid_files, key=get_natural_sort_key):
                sample_id = os.path.splitext(f_name)[0].upper() 
                file_path = os.path.join(root, f_name)
                y_val, x_val = extract_value_at_target_row(file_path, target_row)
                
                if y_val is not None:
                    data_list.append({'Date': date_val, 'Time Folders': folder_name, 'Sample ID': sample_id, 
                                     'Parameter': 'Intensity', 'Measured Value': y_val})
                    data_list.append({'Date': date_val, 'Time Folders': folder_name, 'Sample ID': sample_id, 
                                     'Parameter': 'Time / Wavelength', 'Measured Value': x_val})

    if not data_list:
        print(f"[ WARNING ] No valid data found at row {target_row}.")
        return

    print("\n[ EXPORT ] Generating Excel reports with Statistical Summary...")
    save_professional_excel(pd.DataFrame(data_list), f"Lab_Report_TargetRow_{target_row}.xlsx")
    print("\n[ DONE ] Process completed successfully.")

if __name__ == "__main__":
    CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
    TARGET_FOLDER = os.path.join(CURRENT_DIR, "KSL-202511-Ken") 
    
    if not os.path.exists(TARGET_FOLDER):
        print(f"[ ERROR ] Directory not found: {TARGET_FOLDER}")
        sys.exit()

    max_lines = scan_and_display_file_statistics(TARGET_FOLDER)
    if max_lines > 0:
        target_row = get_user_target_row(max_lines)
        process_professional_layout(TARGET_FOLDER, target_row)
    else:
        print("[ ERROR ] No valid data files detected.")