import os
import sys
import pandas as pd
from tqdm import tqdm
import re
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side

def get_natural_sort_key(s):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', str(s))]

def get_folder_numeric_value(folder_name):
    match = re.search(r'\d+(\.\d+)?', folder_name)
    return float(match.group()) if match else 0.0

def get_folder_status(folder_name):
    return "Diluted" if "dilute" in folder_name.lower() and "not" not in folder_name.lower() else "Undiluted"

# ==========================================
# Terminal Interface & Deep Scan Module
# ==========================================
def scan_and_display_file_statistics(root_folder):
    print("\n[ SYSTEM ] Scanning data directory...")
    max_lines_overall = 0
    file_stats = []
    
    try:
        l1_folders = [f for f in os.listdir(root_folder) if os.path.isdir(os.path.join(root_folder, f)) and f != "Pipeline_Output"]
        l1_folders.sort(key=get_folder_numeric_value)
    except Exception as e:
        print(f"[ ERROR ] Cannot access directory: {e}")
        return 0

    for l1_folder in l1_folders:
        conc = get_folder_numeric_value(l1_folder)
        status = get_folder_status(l1_folder)
        l1_path = os.path.join(root_folder, l1_folder)
        
        for l2_folder in [f for f in os.listdir(l1_path) if os.path.isdir(os.path.join(l1_path, f))]:
            if l2_folder.lower().strip() in ['air', 'di']: continue
            l2_path = os.path.join(l1_path, l2_folder)
            files = sorted([f for f in os.listdir(l2_path) if f.endswith(('.csv', '.txt'))], key=get_natural_sort_key)
            
            sample_files_info = []
            
            for f_name in files:
                file_path = os.path.join(l2_path, f_name)
                try:
                    lines = 0
                    with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                        for line in file:
                            lines += 1
                            if '[EndOfFile]' in line: break
                                
                    if lines > max_lines_overall:
                        max_lines_overall = lines
                        
                    sample_files_info.append({'name': f_name, 'lines': lines})
                except: continue

            if not sample_files_info:
                continue
                
            unique_line_counts = set(info['lines'] for info in sample_files_info)
            
            if len(unique_line_counts) == 1:
                lines = unique_line_counts.pop()
                file_count = len(sample_files_info)
                file_stats.append(f"  |-- [CONC: {conc:<5.2f} | {status.upper():<9}] {l2_folder:<8} {lines} Rows ({file_count} files grouped)")
            else:
                for info in sample_files_info:
                    file_stats.append(f"  |-- [CONC: {conc:<5.2f} | {status.upper():<9}] {l2_folder:<8} -> {info['name']:<25} : {info['lines']} Rows")

    if file_stats:
        print("[ INFO ] File structure analysis complete:")
        for stat in file_stats:
            print(stat)
            
    return max_lines_overall

def get_user_start_row(max_lines):
    print("-" * 70)
    print(f"[ INFO ] Scan complete. Maximum row count is {max_lines}.")
    print("[ INFO ] The program will process data from your starting row to the end of each file.")
    print("[ INFO ] Files with fewer rows than the starting point will be safely skipped.")
    print("[ INFO ] Press [ENTER] without typing anything to cancel and exit.")
    
    while True:
        raw_input = input(f"[ INPUT ] Please enter the starting row (1 - {max_lines}): ").strip()
        
        if not raw_input:
            print("[ CANCEL ] Operation cancelled by user. Exiting program...")
            sys.exit(0)

        try:
            start_row = int(raw_input) 
            
            if start_row <= 0:
                print("[ ERROR ] The value must be greater than zero. Please try again.")
                continue
            if start_row > max_lines:
                print(f"[ ERROR ] The value exceeds the maximum row count ({max_lines}). Please try again.")
                continue
            break 
        except ValueError:
            print("[ ERROR ] Invalid input. Please enter a whole number.")

    return start_row

# ==========================================
# Data Extraction Module
# ==========================================
def extract_peak_and_time_from_file(file_path, start_row):
    max_peak, peak_time = -float('inf'), None
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            for line_idx, line in enumerate(f, start=1):
                if line_idx < start_row: continue
                
                line = line.strip()
                if '[EndOfFile]' in line: break
                if not line or '[Data]' in line: continue

                parts = line.replace(',', ' ').split()
                if len(parts) >= 2:
                    try:
                        x_val, y_val = parts[0], float(parts[1]) 
                        if y_val > max_peak: 
                            max_peak, peak_time = y_val, x_val
                    except ValueError: continue
                        
        return (max_peak, peak_time) if max_peak != -float('inf') else (None, None)
    except: return (None, None)

# ==========================================
# Output Generation Module (Color Magic)
# ==========================================
def save_professional_excel(df_data, file_name):
    if df_data.empty: return

    df_pivot = df_data.pivot_table(
        index='Measurement No.',
        columns=['Conc_Numeric', 'Sample ID', 'Parameter'],
        values='Measured Value',
        aggfunc='first'
    )
    
    df_pivot = df_pivot.sort_index(axis=1, level=[0, 1])

    new_columns = []
    unique_concs = []
    for conc, sample, param in df_pivot.columns:
        conc_str = f"Concentration: {conc:.2f}"
        new_columns.append((conc_str, sample, param))
        if conc_str not in unique_concs:
            unique_concs.append(conc_str)
    
    df_pivot.columns = pd.MultiIndex.from_tuples(new_columns, names=['Concentration', 'Sample ID', 'Parameter'])

    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(base_dir, file_name)
    
    pastel_colors = [
        "DDEBF7", "E2EFDA", "FFF2CC", "FCE4D6", "E4DFEC", "D9D2E9", "F4CCCC"
    ]
    
    conc_color_map = {}
    for idx, conc_str in enumerate(unique_concs):
        conc_color_map[conc_str] = pastel_colors[idx % len(pastel_colors)]

    try:
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df_pivot.to_excel(writer, sheet_name='Analytical Summary')
            ws = writer.sheets['Analytical Summary']
            
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                 top=Side(style='thin'), bottom=Side(style='thin'))
            
            index_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

            for i, col in enumerate(ws.columns, 1):
                max_len = 0
                
                if i == 1:
                    col_fill = index_fill
                else:
                    conc_str = df_pivot.columns[i - 2][0] 
                    hex_color = conc_color_map[conc_str]
                    col_fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

                for cell in col:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = thin_border
                    
                    if i == 1: 
                        if cell.row <= 3: cell.font = Font(bold=True)
                        cell.fill = index_fill
                    else: 
                        cell.fill = col_fill 
                        if cell.row <= 3: 
                            cell.font = Font(bold=True, color="002060") 
                            
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                        
                ws.column_dimensions[get_column_letter(i)].width = max_len + 3
                
        print(f"[ OK ] Export successful: {file_name}")
    except PermissionError:
        print(f"[ ERROR ] Permission denied. Please ensure the Excel file ({file_name}) is closed before running.")

def process_professional_layout(root_folder, start_row):
    print(f"\n[ EXECUTE ] Extracting data starting from row {start_row}...")
    data_list = []
    
    try:
        l1_folders = [f for f in os.listdir(root_folder) if os.path.isdir(os.path.join(root_folder, f)) and f != "Pipeline_Output"]
    except Exception as e: 
        print(f"[ ERROR ] Cannot access directory: {e}")
        return

    for l1_folder in l1_folders:
        conc = get_folder_numeric_value(l1_folder)
        status = get_folder_status(l1_folder)
        l1_path = os.path.join(root_folder, l1_folder)
        
        for l2_folder in [f for f in os.listdir(l1_path) if os.path.isdir(os.path.join(l1_path, f))]:
            if l2_folder.lower().strip() in ['air', 'di']: continue
            
            sample_id = l2_folder.strip().upper()
            
            l2_path = os.path.join(l1_path, l2_folder)
            files = sorted([f for f in os.listdir(l2_path) if f.endswith(('.csv', '.txt'))], key=get_natural_sort_key)
            
            for idx, f_name in enumerate(tqdm(files, desc=f"[ PROCESS ] {l1_folder}/{l2_folder}", leave=False)):
                p_val, p_time = extract_peak_and_time_from_file(os.path.join(l2_path, f_name), start_row)
                
                if p_val is not None:
                    data_list.append({'Condition': status, 'Conc_Numeric': conc, 'Sample ID': sample_id, 
                                     'Measurement No.': idx + 1, 'Parameter': 'Max Intensity', 'Measured Value': p_val})
                    data_list.append({'Condition': status, 'Conc_Numeric': conc, 'Sample ID': sample_id, 
                                     'Measurement No.': idx + 1, 'Parameter': 'Time / Wavelength', 'Measured Value': p_time})

    df_master = pd.DataFrame(data_list)
    if df_master.empty: 
        print(f"[ WARNING ] No data found after row {start_row}.")
        return

    print("\n[ EXPORT ] Generating Excel reports...")
    
    file_diluted = f"Lab_Report_DILUTED_FromRow_{start_row}.xlsx"
    file_undiluted = f"Lab_Report_UNDILUTED_FromRow_{start_row}.xlsx"
    
    save_professional_excel(df_master[df_master['Condition'] == 'Diluted'], file_diluted)
    save_professional_excel(df_master[df_master['Condition'] == 'Undiluted'], file_undiluted)
    
    print("\n[ DONE ] Data extraction and export completed successfully.")

# ==========================================
# Main Execution Protocol
# ==========================================
if __name__ == "__main__":
    CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
    TARGET_FOLDER = os.path.join(CURRENT_DIR, "20260422_KSL_avg50_IT0.65")
    
    if not os.path.exists(TARGET_FOLDER):
        print(f"[ ERROR ] Target folder not found: {TARGET_FOLDER}")
        sys.exit()

    max_lines = scan_and_display_file_statistics(TARGET_FOLDER)
    if max_lines == 0:
        print("[ ERROR ] No valid data files (.csv or .txt) found in the target folder.")
        sys.exit()

    start_row = get_user_start_row(max_lines)
    process_professional_layout(TARGET_FOLDER, start_row)